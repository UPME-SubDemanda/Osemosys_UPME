from __future__ import annotations

import os
import tempfile
from openpyxl import Workbook, load_workbook

from app.services.scenario_export_service import (
    SAND_YEARS,
    _param_is_sand_compatible,
    _sand_row_out,
    _write_sand_rows_from_proxy,
    export_scenario_to_excel_file,
)


def _year_column_index(year: int) -> int:
    return 10 + SAND_YEARS.index(year)


def _make_row(
    *,
    param: str,
    region: str = "R1",
    tech: str = "PWRSOL",
    fuel: str = "",
    emission: str = "",
    timeslice: str = "",
    mode=None,
    storage: str = "",
    year: int | None = 2025,
    value: float = 1.0,
):
    return (
        param,
        region,
        tech,
        fuel,
        emission,
        timeslice,
        mode,
        None,
        None,
        None,
        storage,
        None,
        year,
        value,
    )


class _FakeProxy:
    def __init__(self, rows: list[tuple]) -> None:
        self._rows = rows

    def yield_per(self, _chunk_size: int):
        yield from self._rows


def test_sand_row_out_maps_years_to_columns() -> None:
    key = ("CapitalCost", "R1", "PWRSOL", "", "1", "", "", "", "")
    row = _sand_row_out(key, {2025: 10.0, 2026: 20.0, None: 5.0})
    assert row[0] == "CapitalCost"
    assert row[9] == 5.0
    assert row[_year_column_index(2025)] == 10.0
    assert row[_year_column_index(2026)] == 20.0


def test_param_is_sand_compatible_rejects_udc() -> None:
    assert _param_is_sand_compatible("CapitalCost") is True
    assert _param_is_sand_compatible("UDCMultiplierTotalCapacity") is False


def test_write_sand_rows_groups_multiple_years() -> None:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title="Parameters")
    rows = [
        _make_row(param="CapitalCost", year=2025, value=10.0),
        _make_row(param="CapitalCost", year=2026, value=20.0),
        _make_row(param="CapitalCost", year=2027, value=30.0),
    ]
    source_rows, output_rows = _write_sand_rows_from_proxy(_FakeProxy(rows), ws)
    assert source_rows == 3
    assert output_rows == 1

    key = ("CapitalCost", "R1", "PWRSOL", "", "", "", "", "", "")
    grouped = _sand_row_out(
        key,
        {2025: 10.0, 2026: 20.0, 2027: 30.0},
    )
    assert grouped[_year_column_index(2025)] == 10.0
    assert grouped[_year_column_index(2026)] == 20.0
    assert grouped[_year_column_index(2027)] == 30.0


def test_write_sand_rows_skips_udc_parameters() -> None:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title="Parameters")
    rows = [
        _make_row(param="CapitalCost", year=2025, value=5.0),
        _make_row(param="UDCMultiplierTotalCapacity", year=2025, value=99.0),
    ]
    source_rows, output_rows = _write_sand_rows_from_proxy(_FakeProxy(rows), ws)
    assert source_rows == 1
    assert output_rows == 1


def test_export_scenario_to_excel_file_writes_headers(db_session, monkeypatch) -> None:
    def fake_execute(_query, _params):
        return _FakeProxy([_make_row(param="CapitalCost", year=2025, value=7.0)])

    monkeypatch.setattr(db_session, "execute", fake_execute)

    fd, tmp_path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    try:
        export_scenario_to_excel_file(
            db_session,
            scenario_id=1,
            scenario_name="Test",
            output_path=tmp_path,
        )
        wb = load_workbook(tmp_path, read_only=True)
        try:
            rows = list(wb["Parameters"].iter_rows(values_only=True))
        finally:
            wb.close()
        assert rows[0][0] == "Parameter"
        assert len(rows) == 2
        assert rows[1][0] == "CapitalCost"
        assert rows[1][_year_column_index(2025)] == 7.0
    finally:
        if os.path.exists(tmp_path):
            os.unlink(tmp_path)
