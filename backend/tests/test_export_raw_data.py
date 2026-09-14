from __future__ import annotations

import os
import tempfile
from unittest.mock import patch

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from app.api.v1.visualizations import export_raw_data
from app.models import OsemosysOutputParamValue, SimulationJob
from app.services.simulation_results_export_service import export_raw_data_to_excel_file

from factories import create_scenario, create_user


def _create_succeeded_job(db_session, *, owner, scenario_name: str = "Export test") -> SimulationJob:
    scenario = create_scenario(db_session, name=scenario_name, owner=owner.username)
    job = SimulationJob(
        user_id=owner.id,
        scenario_id=scenario.id,
        solver_name="highs",
        status="SUCCEEDED",
        progress=100.0,
    )
    db_session.add(job)
    db_session.flush()
    return job


def test_export_raw_data_to_excel_file_writes_expected_rows(db_session) -> None:
    owner = create_user(db_session, username="export-owner")
    job = _create_succeeded_job(db_session, owner=owner)
    db_session.add_all(
        [
            OsemosysOutputParamValue(
                id_simulation_job=job.id,
                variable_name="RateOfActivity",
                technology_name="AN_PWRSOL",
                fuel_name="ELC",
                year=2030,
                value=12.5,
                index_json={"REGION": "AN"},
            ),
            OsemosysOutputParamValue(
                id_simulation_job=job.id,
                variable_name="TotalDiscountedCost",
                value=999.0,
            ),
        ]
    )
    db_session.commit()

    fd, tmp_path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    try:
        row_count = export_raw_data_to_excel_file(
            db_session, job_id=job.id, output_path=tmp_path
        )
        assert row_count == 2
        wb = load_workbook(tmp_path, read_only=True)
        try:
            ws = wb["Raw Data"]
            rows = list(ws.iter_rows(values_only=True))
        finally:
            wb.close()
        assert rows[0] == (
            "VariableName",
            "Technology",
            "Fuel",
            "Emission",
            "Year",
            "Value",
            "IndexJSON",
        )
        assert rows[1][0] == "RateOfActivity"
        assert rows[1][1] == "AN_PWRSOL"
        assert rows[1][5] == 12.5
        assert rows[2][0] == "TotalDiscountedCost"
        assert rows[2][5] == 999.0
    finally:
        if os.path.exists(tmp_path):
            os.unlink(tmp_path)


def test_export_raw_data_to_excel_file_raises_404_when_empty(db_session) -> None:
    owner = create_user(db_session, username="export-empty")
    job = _create_succeeded_job(db_session, owner=owner)
    db_session.commit()

    fd, tmp_path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    try:
        with pytest.raises(HTTPException) as exc:
            export_raw_data_to_excel_file(
                db_session, job_id=job.id, output_path=tmp_path
            )
        assert exc.value.status_code == 404
    finally:
        if os.path.exists(tmp_path):
            os.unlink(tmp_path)


def test_export_raw_data_to_excel_file_streams_many_rows(db_session) -> None:
    owner = create_user(db_session, username="export-bulk")
    job = _create_succeeded_job(db_session, owner=owner)
    batch_size = 5_000
    db_session.add_all(
        OsemosysOutputParamValue(
            id_simulation_job=job.id,
            variable_name="RateOfActivity",
            technology_name=f"TECH_{idx}",
            year=2025 + (idx % 5),
            value=float(idx),
        )
        for idx in range(batch_size)
    )
    db_session.commit()

    fd, tmp_path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    try:
        with patch(
            "app.services.simulation_results_export_service._YIELD_PER",
            500,
        ):
            row_count = export_raw_data_to_excel_file(
                db_session, job_id=job.id, output_path=tmp_path
            )
        assert row_count == batch_size
        wb = load_workbook(tmp_path, read_only=True)
        try:
            ws = wb["Raw Data"]
            data_rows = sum(1 for _ in ws.iter_rows(min_row=2, values_only=True))
            assert data_rows == batch_size
        finally:
            wb.close()
    finally:
        if os.path.exists(tmp_path):
            os.unlink(tmp_path)


def test_export_raw_data_endpoint_returns_file_response(db_session) -> None:
    owner = create_user(db_session, username="export-endpoint")
    job = _create_succeeded_job(db_session, owner=owner, scenario_name="Regional 2TS")
    db_session.add(
        OsemosysOutputParamValue(
            id_simulation_job=job.id,
            variable_name="Dispatch",
            technology_name="CA_PWRSOL",
            value=1.0,
            year=2030,
        )
    )
    db_session.commit()

    job_public = {
        "id": job.id,
        "status": "SUCCEEDED",
        "scenario_name": "Regional 2TS",
    }

    with patch(
        "app.api.v1.visualizations.SimulationService.get_by_id",
        return_value=job_public,
    ):
        response = export_raw_data(
            job_id=job.id,
            db=db_session,
            current_user=owner,
        )

    assert response.status_code == 200
    assert (
        response.media_type
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    disposition = response.headers.get("content-disposition", "")
    assert "Resultados_Crudos_Regional" in disposition
    assert disposition.endswith(".xlsx")
    assert os.path.exists(response.path)
